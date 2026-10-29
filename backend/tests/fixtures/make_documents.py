"""Build the document fixtures for the parser integration tests (F-02).

The legal excerpts are pages 53-54 of Regulation (EU) 2024/1689 (the AI Act) in English
and Polish: the end of Article 5, all of Article 6 and the start of Article 7. Source:
the Publications Office of the European Union, reused under Commission Decision
2011/833/EU (© European Union, https://eur-lex.europa.eu/eli/reg/2024/1689/oj).

Everything else is generated here. Run from `backend/`; it needs the network once, and
pikepdf only for the password-protected file:

    uv run --with pikepdf python tests/fixtures/make_documents.py
"""

import random
import urllib.request
from pathlib import Path

import pikepdf
import pypdfium2 as pdfium
from docx import Document as WordDocument
from PIL import Image, ImageFilter

OUT = Path(__file__).parent / "documents"
SOURCE = "https://publications.europa.eu/resource/celex/32024R1689"
LANGUAGES = {"en": "eng", "pl": "pol"}
# Zero-based page indexes of Articles 5 (end) to 7 (start) in the Official Journal PDF.
EXCERPT_PAGES = [52, 53]
SCAN_DPI = 200


def fetch_regulation(language: str) -> pdfium.PdfDocument:
    request = urllib.request.Request(
        SOURCE, headers={"Accept": "application/pdf", "Accept-Language": language}
    )
    with urllib.request.urlopen(request, timeout=60) as response:  # noqa: S310 (fixed https URL)
        return pdfium.PdfDocument(response.read())


def save_pages(source: pdfium.PdfDocument, pages: list[int], path: Path) -> None:
    excerpt = pdfium.PdfDocument.new()
    excerpt.import_pages(source, pages)
    excerpt.save(path)


def scanned(page: pdfium.PdfPage, seed: int) -> Image.Image:
    """A page as a scanner would see it: grey, slightly rotated, a little noisy."""
    image: Image.Image = page.render(scale=SCAN_DPI / 72).to_pil().convert("L")
    image = image.rotate(0.5, expand=True, fillcolor=255).filter(ImageFilter.GaussianBlur(0.3))
    noise = random.Random(seed)  # noqa: S311 (repeatable noise, not security)
    width, height = image.size
    for _ in range(width * height // 5000):
        x, y = noise.randrange(width), noise.randrange(height)
        image.putpixel((x, y), noise.choice((0, 255)))
    return image


def write_docx(path: Path) -> None:
    document = WordDocument()
    document.add_heading("Data retention policy", level=1)
    document.add_paragraph("This policy sets how long the team keeps records of model use.")
    document.add_heading("Retention periods", level=2)
    document.add_paragraph("Logs are kept for six months and then deleted.")
    table = document.add_table(rows=3, cols=2)
    for row, (record, period) in enumerate(
        [("Record", "Kept for"), ("Access logs", "6 months"), ("Incident reports", "5 years")]
    ):
        table.cell(row, 0).text = record
        table.cell(row, 1).text = period
    document.add_heading("Review", level=2)
    document.add_paragraph("The policy is reviewed every year by the compliance officer.")
    document.save(str(path))


def write_blank_pdf(path: Path, pages: int) -> None:
    blank = pdfium.PdfDocument.new()
    for _ in range(pages):
        blank.new_page(595, 842)
    blank.save(path)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    regulations = {short: fetch_regulation(code) for short, code in LANGUAGES.items()}
    for short, regulation in regulations.items():
        save_pages(regulation, EXCERPT_PAGES, OUT / f"ai-act-art6-{short}.pdf")

    # Article 6 starts on the first excerpt page.
    english_page = regulations["en"][EXCERPT_PAGES[0]]
    scan = scanned(english_page, seed=1)
    scan.save(OUT / "ai-act-art6-en-scan.pdf", resolution=SCAN_DPI)
    scan.save(OUT / "ai-act-art6-en-photo.jpg", quality=80)
    scanned(regulations["pl"][EXCERPT_PAGES[0]], seed=2).save(OUT / "ai-act-art6-pl-photo.png")

    write_docx(OUT / "retention-policy.docx")

    with pikepdf.open(OUT / "ai-act-art6-en.pdf") as source:
        source.save(OUT / "locked.pdf", encryption=pikepdf.Encryption(owner="owner", user="user"))
    corrupt = (OUT / "ai-act-art6-en.pdf").read_bytes()[:2048]
    (OUT / "corrupt.pdf").write_bytes(corrupt)
    write_blank_pdf(OUT / "blank.pdf", pages=1)
    write_blank_pdf(OUT / "too-long.pdf", pages=501)

    sizes = {path.name: path.stat().st_size for path in sorted(OUT.iterdir())}
    print("\n".join(f"{size:>9}  {name}" for name, size in sizes.items()))


if __name__ == "__main__":
    main()
